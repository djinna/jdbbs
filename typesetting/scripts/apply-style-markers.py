#!/usr/bin/env python3
"""Apply [[style]] markers in a .docx — a pre-pass for authors without custom styles.

Authors drafting in Google Docs have no custom paragraph styles, so their intent
arrives as ad-hoc text markers. This script makes the markers a convention the
factory understands:

    [[style:computer text]]Hello, would you like to hear a TCP joke?
    [[code block]] ...          (the "style:" prefix is optional)
    ... last line [[/code block]]   (closer: name optional, need not match)
    ... last line [[/]]  or  [[end]]

Rules (keep in step with detect_style_markers() in detect-edge-cases.py, which
imports this file):

  * An opening marker is recognised only at the START of a paragraph; a closing
    marker only at the END of a paragraph.
  * Opener with no closer -> that paragraph only. Opener with a closer on a
    later (or the same) paragraph -> every paragraph from opener to closer
    inclusive. Nested/overlapping: a closer closes the most recent unclosed
    opener (a named closer prefers an opener of the same name); later ranges
    win where they overlap.
  * The name is matched case-insensitively, ignoring spaces, hyphens and
    underscores, against the factory paragraph styles, a list of aliases
    (computer text / code / mono -> Code Block, quote -> Block Quote, poem ->
    Verse, ...), and the custom styles declared on the transmittal.
  * Resolved: paragraph.style is set to the style (added to the document if
    missing) and the marker text is removed. Unresolved: the marker is left in
    place so it shows in the proof, and it is reported.

Usage:
    apply-style-markers.py IN.docx OUT.docx [--declared-styles declared.json] [--report report.json]

IN may equal OUT. Exit status is 0 whenever the docx is readable.
"""

import argparse
import json
import os
import re
import shutil
import sys
import tempfile
from typing import Dict, List, Optional, Tuple

# --- marker grammar ---------------------------------------------------------

OPEN_RE = re.compile(r"^\s*\[\[\s*(?:style\s*:\s*)?([^\]/][^\]]*?)\s*\]\]\s*")
CLOSE_RE = re.compile(r"\s*\[\[\s*(?:/\s*([^\]]*?)|end)\s*\]\]\s*$", re.IGNORECASE)

# --- style resolution -------------------------------------------------------

FACTORY_PARAGRAPH_STYLES = [
    "Normal", "First Paragraph", "Heading 1", "Heading 2", "Heading 3",
    "Block Quote", "Epigraph", "Verse", "Code Block", "Section Break",
    "Copyright", "Signature", "Glossary Entry",
]

# alias (normalised) -> factory style
ALIASES = {
    "computertext": "Code Block", "code": "Code Block", "terminal": "Code Block",
    "monospace": "Code Block", "mono": "Code Block",
    "quote": "Block Quote", "blockquote": "Block Quote", "extract": "Block Quote",
    "poem": "Verse", "poetry": "Verse",
    "epi": "Epigraph",
    "firstpara": "First Paragraph", "noindent": "First Paragraph",
    "break": "Section Break", "scenebreak": "Section Break", "sectionbreak": "Section Break",
    "h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3", "chapter": "Heading 1",
}

# Text for a Section Break paragraph left empty by a bare [[break]] marker
# (matches SECTION_BREAK_CHARS["breve"] in generate-word-template.py).
SECTION_BREAK_PLACEHOLDER = "\u02D8"

# Author-facing list for the "not a factory style" message.
SUGGESTED_NAMES = "code block, block quote, epigraph, verse, first paragraph, section break"


def normalize(name: str) -> str:
    """Same normalisation as the Lua filter: lower, drop spaces/-/_.

    Marker brackets typed around a declared name ("[[commentary]]") are
    stripped too, so it matches the marker [[commentary]] in the text.
    """
    name = (name or "").strip()
    while True:
        m = re.match(r"^\[\[\s*(.*?)\s*\]\]$|^\[\s*(.*?)\s*\]$", name)
        if not m:
            break
        name = (m.group(1) if m.group(1) is not None else m.group(2) or "").strip()
    return re.sub(r"[\s\-_]", "", name.lower())


class StyleResolver:
    """Resolve a marker name to (Word style name, via) or (None, None)."""

    def __init__(self, declared_styles: Optional[List] = None):
        self.factory = {normalize(n): n for n in FACTORY_PARAGRAPH_STYLES}
        self.declared: Dict[str, str] = {}
        for item in declared_styles or []:
            if isinstance(item, str):
                name, word = item, item
            elif isinstance(item, dict):
                name = (item.get("name") or "").strip()
                word = (item.get("word_style") or "").strip() or name
                if (item.get("type") or "paragraph") == "character":
                    continue
            else:
                continue
            if not word:
                continue
            # The Word style name is what pandoc emits and the Lua filter maps on.
            for key in (name, word):
                if normalize(key):
                    self.declared.setdefault(normalize(key), word)

    def resolve(self, marker: str) -> Tuple[Optional[str], Optional[str]]:
        key = normalize(marker)
        if not key or key == "end":
            return None, None
        if key in self.factory:
            return self.factory[key], "factory"
        if key in ALIASES:
            return ALIASES[key], "alias"
        if key in self.declared:
            return self.declared[key], "declared"
        return None, None


def load_declared_styles(path: str) -> List:
    if not path or not os.path.exists(path):
        return []
    try:
        with open(path, encoding="utf-8") as fh:
            data = json.load(fh)
    except (OSError, ValueError):
        return []
    return data if isinstance(data, list) else []


# --- scanning -----------------------------------------------------------------

def paragraph_runs_text(para) -> str:
    return "".join(r.text or "" for r in para.runs)


def scan_markers(paragraphs, resolver: StyleResolver) -> List[Dict]:
    """Find every opener and pair it with a closer. Returns one dict per opener:

    {marker, resolved, via, start, end, paragraphs, open_len, close_index}
    (start/end are 0-based here; open_len is the number of characters to strip
    from the front of the opening paragraph; close_index is the 0-based
    paragraph carrying the closer, or None.)
    """
    found: List[Dict] = []
    open_stack: List[Dict] = []  # unclosed openers, in document order
    for i, para in enumerate(paragraphs):
        text = paragraph_runs_text(para)
        m = OPEN_RE.match(text)
        if m and normalize(m.group(1)) not in ("", "end"):
            name = m.group(1).strip()
            resolved, via = resolver.resolve(name)
            entry = {
                "marker": name, "resolved": resolved, "via": via,
                "start": i, "end": i, "open_len": m.end(), "close_index": None,
            }
            found.append(entry)
            open_stack.append(entry)
            text_after_open = text[m.end():]
        else:
            text_after_open = text
        c = CLOSE_RE.search(text_after_open)
        if c and open_stack:
            cname = normalize(c.group(1) or "")
            target = None
            if cname:
                for entry in reversed(open_stack):
                    if normalize(entry["marker"]) == cname or normalize(entry["resolved"] or "") == cname \
                            or normalize(ALIASES.get(cname, "")) == normalize(entry["resolved"] or ""):
                        target = entry
                        break
            if target is None:
                target = open_stack[-1]
            open_stack.remove(target)
            target["end"] = i
            target["close_index"] = i
            # characters to strip from the END of paragraph i
            target["close_strip"] = len(text_after_open) - c.start()
    for entry in found:
        entry["paragraphs"] = entry["end"] - entry["start"] + 1
    return found


def report_entries(found: List[Dict]) -> List[Dict]:
    """Public shape, 1-based paragraph numbers (matching Inspect)."""
    return [{
        "marker": e["marker"], "resolved": e["resolved"], "via": e["via"],
        "start": e["start"] + 1, "end": e["end"] + 1, "paragraphs": e["paragraphs"],
    } for e in found]


# --- applying -----------------------------------------------------------------

def strip_from_start(para, n: int) -> None:
    """Remove the first n characters of the paragraph's run text (markers may span runs)."""
    for run in para.runs:
        if n <= 0:
            break
        t = run.text or ""
        if len(t) <= n:
            n -= len(t)
            run.text = ""
        else:
            run.text = t[n:]
            n = 0


def strip_from_end(para, n: int) -> None:
    for run in reversed(para.runs):
        if n <= 0:
            break
        t = run.text or ""
        if len(t) <= n:
            n -= len(t)
            run.text = ""
        else:
            run.text = t[:len(t) - n]
            n = 0


def find_style(doc, name: str):
    """Paragraph style whose name matches `name` after normalisation, or None.
    (Google Docs exports 'normal', Word 'Normal'; the Lua filter treats them alike.)"""
    from docx.enum.style import WD_STYLE_TYPE
    want = normalize(name)
    for style in doc.styles:
        if style.type == WD_STYLE_TYPE.PARAGRAPH and normalize(style.name) == want:
            return style
    return None


def ensure_style(doc, name: str):
    from docx.enum.style import WD_STYLE_TYPE
    style = find_style(doc, name)
    if style is not None:
        return style
    style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    base = find_style(doc, "Normal")
    if base is not None:
        style.base_style = base
    return style


def clear_direct_indent(para) -> None:
    """Drop the paragraph's own <w:ind>. The style now owns the layout; left over,
    pandoc reads a direct left indent as a block quote and nests it inside the Div."""
    pPr = para._p.pPr
    if pPr is None:
        return
    for ind in pPr.findall(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ind"):
        pPr.remove(ind)


def apply_markers(doc, resolver: StyleResolver) -> List[Dict]:
    paragraphs = doc.paragraphs
    found = scan_markers(paragraphs, resolver)
    # Later ranges win where they overlap: apply in document order.
    for e in found:
        if not e["resolved"]:
            continue
        style = ensure_style(doc, e["resolved"])
        for i in range(e["start"], e["end"] + 1):
            paragraphs[i].style = style
            clear_direct_indent(paragraphs[i])
    # Strip marker text; closers first so opener offsets on the same paragraph stay valid.
    for e in found:
        if not e["resolved"]:
            continue
        if e["close_index"] is not None:
            strip_from_end(paragraphs[e["close_index"]], e["close_strip"])
    for e in found:
        if e["resolved"]:
            strip_from_start(paragraphs[e["start"]], e["open_len"])
    # A bare "[[break]]" paragraph is empty once the marker is gone, and pandoc
    # drops empty paragraphs — so the section/stanza break vanished from the
    # book (Devotion, 2026-09-21). Give it the same ornament text the Word
    # template's Section Break paragraphs carry; the Typst filter ignores the
    # content of a section break anyway.
    for e in found:
        if e["resolved"] == "Section Break":
            for i in range(e["start"], e["end"] + 1):
                if not paragraphs[i].text.strip():
                    paragraphs[i].add_run(SECTION_BREAK_PLACEHOLDER)
    return found


def main() -> int:
    ap = argparse.ArgumentParser(description="Apply [[style]] markers in a .docx")
    ap.add_argument("input")
    ap.add_argument("output")
    ap.add_argument("--declared-styles", default="", help="JSON array of declared custom styles")
    ap.add_argument("--report", default="", help="write a JSON report of markers here")
    args = ap.parse_args()

    from docx import Document
    doc = Document(args.input)
    resolver = StyleResolver(load_declared_styles(args.declared_styles))
    found = apply_markers(doc, resolver)
    report = report_entries(found)

    out_dir = os.path.dirname(os.path.abspath(args.output)) or "."
    fd, tmp = tempfile.mkstemp(prefix=".style-markers-", suffix=".docx", dir=out_dir)
    os.close(fd)
    try:
        doc.save(tmp)
        shutil.move(tmp, args.output)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)

    if args.report:
        with open(args.report, "w", encoding="utf-8") as fh:
            json.dump(report, fh, indent=2)

    resolved = [r for r in report if r["resolved"]]
    unresolved = [r for r in report if not r["resolved"]]
    parts = [f"{len(resolved)} resolved"]
    if resolved:
        parts.append(", ".join(f"[[{r['marker']}]]\u2192{r['resolved']} \u00b6{r['start']}" +
                               (f"\u2013{r['end']}" if r['end'] != r['start'] else "") for r in resolved))
    parts.append(f"{len(unresolved)} unresolved")
    if unresolved:
        parts.append(", ".join(f"[[{r['marker']}]] \u00b6{r['start']}" for r in unresolved))
    print("style markers: " + "; ".join(parts), file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
